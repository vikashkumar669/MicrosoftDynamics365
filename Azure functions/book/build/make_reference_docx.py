"""Builds build/reference.docx: the Word styles Pandoc uses for the DOCX
manuscript. Run once after changing the design system:

    pandoc -o /tmp/default.docx --print-default-data-file reference.docx
    python3 make_reference_docx.py /tmp/default.docx reference.docx
"""

import sys

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = "0B1F3A"
INK = "13233D"
AZURE = "0F6CBD"
SKY = "E8F2FB"
SLATE = "4A5568"
SANS = "Inter"
SERIF = "Source Serif 4"
MONO = "Consolas"  # ships with Windows and Office; print PDF uses Cascadia Code


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_font(style, name, size=None, color=None, bold=None, italic=None):
    font = style.font
    font.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    if size is not None:
        font.size = Pt(size)
    if color is not None:
        font.color.rgb = rgb(color)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic


def set_spacing(style, before=0, after=6, line=None):
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = line


def shade(style, fill, border_color=None, border_side="left"):
    ppr = style.element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)
    if border_color:
        bdr = OxmlElement("w:pBdr")
        side = OxmlElement("w:" + border_side)
        side.set(qn("w:val"), "single")
        side.set(qn("w:sz"), "18")
        side.set(qn("w:space"), "6")
        side.set(qn("w:color"), border_color)
        bdr.append(side)
        ppr.append(bdr)
    ind = style.paragraph_format
    ind.left_indent = Inches(0.08)
    ind.right_indent = Inches(0.08)


class StyleLookup:
    """Finds styles by their stored name. python-docx translates built-in
    names (\"Heading 1\" -> \"heading 1\"), which misses Pandoc's file."""

    def __init__(self, doc):
        self.doc = doc

    def __getitem__(self, name):
        for style in self.doc.styles:
            if style.name == name:
                return style
        raise KeyError(name)


def para_style(doc, name, base="Normal"):
    lookup = StyleLookup(doc)
    try:
        style = lookup[name]
    except KeyError:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = lookup[base]
    style.quick_style = True
    return style


def add_page_number_footer(section):
    footer = section.footer
    paragraph = footer.paragraphs[0]
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), kind)
            run._r.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = text
            run._r.append(instr)
    run.font.name = SANS
    run.font.size = Pt(8.5)
    run.font.color.rgb = rgb(AZURE)


def main(src, dst):
    doc = Document(src)

    # Page: 7.5 x 9.25 in, mirrored margins (inside 0.875 in).
    for section in doc.sections:
        section.page_width = Inches(7.5)
        section.page_height = Inches(9.25)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(0.875)
        section.right_margin = Inches(0.65)
        section.gutter = Inches(0)
        add_page_number_footer(section)
    settings = doc.settings.element
    mirror = OxmlElement("w:mirrorMargins")
    settings.insert(0, mirror)

    s = StyleLookup(doc)
    set_font(s["Normal"], SERIF, 10.5, INK)
    set_spacing(s["Normal"], 0, 6, 1.2)
    for name in ("Body Text", "First Paragraph", "Compact"):
        try:
            set_font(s[name], SERIF, 10.5, INK)
        except KeyError:
            pass

    set_font(s["Title"], SANS, 28, NAVY, bold=True)
    set_spacing(s["Title"], 0, 4)
    set_font(s["Subtitle"], SANS, 11, AZURE, bold=True)
    set_spacing(s["Subtitle"], 0, 18)
    set_font(s["Heading 1"], SANS, 22, NAVY, bold=True)
    set_spacing(s["Heading 1"], 24, 8)
    set_font(s["Heading 2"], SANS, 15, NAVY, bold=True)
    set_spacing(s["Heading 2"], 18, 6)
    set_font(s["Heading 3"], SANS, 11.5, AZURE, bold=True)
    set_spacing(s["Heading 3"], 12, 4)

    # Code: light panel in the editable manuscript (print PDF is dark).
    code = para_style(doc, "Source Code")
    set_font(code, MONO, 8.5, NAVY)
    set_spacing(code, 0, 0, 1.0)
    shade(code, "F3F7FB", AZURE)
    try:
        set_font(s["Verbatim Char"], MONO, 9, "0B4F8A")
    except KeyError:
        pass

    for name in ("Image Caption", "Table Caption", "Listing Caption"):
        style = para_style(doc, name)
        set_font(style, SANS, 8.5, SLATE)
        set_spacing(style, 2, 8)

    callouts = {
        "Tip": (SKY, AZURE),
        "Note": ("F2F4F7", SLATE),
        "Warning": ("FFF4E0", "9A5100"),
        "Security": ("FDECEC", "B4232A"),
        "Performance": ("E6F5F1", "0E7C66"),
        "Objectives": (SKY, AZURE),
        "Exercise": ("FFFFFF", AZURE),
        "Knowledge Check": ("FFFFFF", AZURE),
        "Summary": (NAVY, NAVY),
    }
    for name, (fill, color) in callouts.items():
        label = para_style(doc, name + " Label")
        set_font(label, SANS, 8,
                 "9CCBFA" if name == "Summary" else color, bold=True)
        label.font.all_caps = True
        set_spacing(label, 8, 2)
        shade(label, fill, color)
        body = para_style(doc, name + " Body")
        set_font(body, SANS, 9.5,
                 "E6EDF5" if name == "Summary" else INK)
        set_spacing(body, 0, 3, 1.15)
        shade(body, fill, color)

    capture = para_style(doc, "Capture Box")
    set_font(capture, SANS, 8.5, SLATE)
    set_spacing(capture, 0, 2)
    shade(capture, "EEF2F7", "8A9BB0")

    opener = {
        "Opener Part": (SANS, 8, AZURE, True),
        "Opener Summary": (SERIF, 11, SLATE, False),
        "Opener Builds Label": (SANS, 8, AZURE, True),
        "Opener Builds": (SANS, 10, NAVY, False),
        "Opener Versions": (SANS, 8, SLATE, False),
    }
    for name, (font, size, color, bold) in opener.items():
        style = para_style(doc, name)
        set_font(style, font, size, color, bold=bold)
        set_spacing(style, 0, 6)
    s["Opener Part"].font.all_caps = True
    s["Opener Builds Label"].font.all_caps = True
    s["Opener Summary"].font.italic = True

    doc.save(dst)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
