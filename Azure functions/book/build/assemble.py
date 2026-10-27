"""Assembles the whole book: one PDF (print layout) and one DOCX.

    python3 book/build/assemble.py

Run book/build/build.sh for every chapter and appendix first; this script
joins their HTML (for the PDF) and their DOCX files (for the manuscript),
adding the cover, title page, copyright page, table of contents and
preface.
"""

import glob
import html
import os
import re
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
BOOK = os.path.dirname(HERE)
CODE = os.path.dirname(BOOK)
OUT = os.path.join(BOOK, "out")

TITLE = "Azure Functions for Microsoft Dynamics 365 &amp; Power Platform"
SUBTITLE = ("A Practical Guide to Integration, Automation, APIs, and "
            "Enterprise Solutions")
TAGLINE = "From Fundamentals to Production-Ready Enterprise Implementations"
AUTHOR = "Vikash Kumar"


def sources():
    chapters = sorted(glob.glob(os.path.join(BOOK, "chapters", "ch*.md")))
    appendices = sorted(glob.glob(os.path.join(BOOK, "chapters", "app*.md")))
    return chapters + appendices


def meta(md):
    text = open(md, encoding="utf-8").read()
    head = text.split("---", 2)[1]
    get = lambda k: (re.search(rf'^{k}:\s*"?(.*?)"?\s*$', head, re.M)
                     or [None, ""])[1]
    return {"chapter": get("chapter"), "title": get("title"),
            "part": get("part"), "appendix": "appendix: true" in head}


def body_of(path, prefix):
    doc = open(path, encoding="utf-8").read()
    body = doc.split("<body>", 1)[1].rsplit("</body>", 1)[0]
    # Make ids unique across chapters.
    body = re.sub(r'id="([^"]+)"', lambda m: f'id="{prefix}-{m.group(1)}"',
                  body)
    body = re.sub(r'href="#([^"]+)"',
                  lambda m: f'href="#{prefix}-{m.group(1)}"', body)
    return body


def versions_table():
    rows = set()
    for proj in glob.glob(os.path.join(CODE, "**", "*.csproj"),
                          recursive=True):
        for m in re.finditer(r'Include="([^"]+)" Version="([^"]+)"',
                             open(proj).read()):
            rows.add(m.groups())
    keep = [r for r in sorted(rows) if not r[0].startswith("System.")]
    items = "; ".join(f"{html.escape(a)} {b}" for a, b in keep)
    return f"<p class='versions'>{items}.</p>"


def cover():
    svg = open(os.path.join(HERE, "cover-art.svg"), encoding="utf-8").read()
    return f"""
<section class="cover">
  <div class="cover-top">
    <p class="cover-kicker">A Practical Guide</p>
    <p class="cover-pill">.NET 10 · Isolated worker · Flex Consumption</p>
  </div>
  <div class="cover-accent"></div>
  <h1 class="cover-title">Azure Functions</h1>
  <p class="cover-sub">for Microsoft Dynamics 365 &amp; Power Platform</p>
  <p class="cover-desc">{SUBTITLE}</p>
  <div class="cover-art">{svg}</div>
  <div class="cover-foot">
    <p class="cover-tag">{TAGLINE}</p>
    <div class="cover-row"><p class="cover-author">{AUTHOR}</p>
    <p class="cover-count">25 chapters · 4 capstone projects</p></div>
  </div>
</section>"""


def title_and_copyright():
    return f"""
<section class="titlepage">
  <p class="tp-title">{TITLE}</p>
  <p class="tp-sub">{SUBTITLE}</p>
  <p class="tp-tag">{TAGLINE}</p>
  <p class="tp-author">{AUTHOR}</p>
</section>
<section class="copyright">
  <p><strong>{TITLE}</strong><br>{SUBTITLE}</p>
  <p>Copyright © 2026 Vikash Kumar. All rights reserved. No part of this
  book's text or figures may be reproduced without the author's written
  permission, except for brief quotations in reviews. The companion source
  code is licensed separately under the MIT licence.</p>
  <p>Self-published by the author. First edition, October 2026.</p>
  <p><strong>Trademarks.</strong> Microsoft, Azure, Dynamics 365, Power
  Platform, Power Apps, Power Automate, Power Pages, Dataverse, Excel and
  Microsoft Entra are trademarks of the Microsoft group of companies. Other
  names may be trademarks of their owners. They are used for identification
  only. This book is independent and is not endorsed by or affiliated with
  Microsoft. Azure architecture icons in diagrams are used under
  Microsoft's terms for architecture diagrams and documentation.</p>
  <p><strong>Fictional content.</strong> Harbourline Group, Harbourline
  Connect, LedgerOne, SignBridge and all people, properties and data in
  this book are fictional.</p>
  <p><strong>Disclaimer.</strong> The information and code are provided
  "as is", without warranty. Cloud services change; verify limits,
  versions and portal paths against the cited Microsoft documentation
  before production use. The author accepts no liability for loss arising
  from the use of this book.</p>
  <p><strong>Technical version statement.</strong> Written and tested
  against .NET 10, Azure Functions runtime 4.x (isolated worker model,
  <code>Azure.Functions.Sdk</code> 1.0.0) on the Flex Consumption plan, and
  Microsoft Dataverse, with documentation checked in October 2026. The
  companion code is compiled and unit-tested in continuous integration with
  these packages:</p>
  {versions_table()}
</section>"""


def toc(entries):
    items = []
    last_part = None
    for e in entries:
        if not e["appendix"] and e["part"] != last_part:
            items.append(f'<li class="toc-part">{html.escape(e["part"])}</li>')
            last_part = e["part"]
        if e["appendix"] and last_part != "Appendices":
            items.append('<li class="toc-part">Appendices</li>')
            last_part = "Appendices"
        label = (e["chapter"] if e["appendix"]
                 else f"{int(e['chapter'])}")
        items.append(f'<li><a href="#{e["anchor"]}"><span class="toc-n">'
                     f'{label}</span> {html.escape(e["title"])}</a></li>')
    return ('<section class="toc"><h2 class="toc-head">Contents</h2>'
            '<ul class="toc-list"><li><a href="#front-preface">'
            '<span class="toc-n"></span> Preface</a></li>'
            + "".join(items) + "</ul></section>")


def preface_html():
    md = os.path.join(BOOK, "front", "preface.md")
    out = subprocess.run(["pandoc", md, "--to", "html5"],
                         capture_output=True, text=True, check=True).stdout
    return f'<section class="preface" id="front-preface">{out}</section>'


def build_pdf(entries):
    parts = []
    for e in entries:
        body = body_of(e["html"], e["anchor"])
        parts.append(f'<div class="chapter-wrap" id="{e["anchor"]}">'
                     f"{body}</div>")
    doc = f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8">
<title>{TITLE}</title>
<meta name="author" content="{AUTHOR}">
<meta name="description" content="{SUBTITLE}">
<link rel="stylesheet" href="book.css">
<link rel="stylesheet" href="front.css">
</head><body>
{cover()}
{title_and_copyright()}
{toc(entries)}
{preface_html()}
{''.join(parts)}
</body></html>"""
    path = os.path.join(OUT, "book.html")
    open(path, "w", encoding="utf-8").write(doc)
    subprocess.run(["cp", os.path.join(HERE, "front.css"), OUT], check=True)
    pdf = os.path.join(OUT, "Azure-Functions-for-Dynamics-365.pdf")
    subprocess.run(["weasyprint", "--base-url", OUT + "/", path, pdf],
                   check=True)
    return pdf


def build_docx(entries):
    from docx import Document
    from docxcompose.composer import Composer

    front_md = os.path.join(OUT, "front.md")
    open(front_md, "w", encoding="utf-8").write(
        f"---\ntitle: \"{html.unescape(TITLE)}\"\n"
        f"subtitle: \"{SUBTITLE}\"\nauthor: \"{AUTHOR}\"\n---\n\n"
        "Copyright © 2026 Vikash Kumar. All rights reserved. "
        "Self-published. Companion code: MIT licence. All companies, "
        "people and data are fictional. Microsoft product names are "
        "trademarks of the Microsoft group of companies; this book is "
        "not endorsed by Microsoft.\n\n"
        + open(os.path.join(BOOK, "front", "preface.md"),
               encoding="utf-8").read())
    front_docx = os.path.join(OUT, "front.docx")
    subprocess.run(["pandoc", front_md, "-o", front_docx,
                    "--reference-doc", os.path.join(HERE, "reference.docx"),
                    "--toc", "--toc-depth=1"], check=True)
    master = Document(front_docx)
    composer = Composer(master)
    for e in entries:
        doc = Document(e["docx"])
        doc.paragraphs[0].insert_paragraph_before().paragraph_format \
            .page_break_before = True
        composer.append(doc)
    path = os.path.join(OUT, "Azure-Functions-for-Dynamics-365.docx")
    composer.save(path)
    return path


def main():
    entries = []
    for md in sources():
        name = os.path.splitext(os.path.basename(md))[0]
        e = meta(md)
        e["html"] = os.path.join(OUT, name + ".html")
        e["docx"] = os.path.join(OUT, name + ".docx")
        e["anchor"] = name
        if not os.path.exists(e["html"]):
            sys.exit(f"Build {name} first (book/build/build.sh).")
        entries.append(e)
    pdf = build_pdf(entries)
    print("PDF:", pdf)
    if "--no-docx" not in sys.argv:
        print("DOCX:", build_docx(entries))


if __name__ == "__main__":
    main()
